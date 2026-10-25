# Importa las hojas «Videoteca» y «Discoteca» de Wishlist.xlsx a public/coleccion.json.
# Solo hace falta una vez: a partir de ahí la colección se edita en la web.
#
#   python scripts/importar-excel.py "C:\Users\adric\Desktop\Wishlist.xlsx"
#
# Si public/coleccion.json ya existe no se sobrescribe, salvo con --forzar.
import json
import sys
import uuid
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / 'public' / 'coleccion.json'

FORMATOS = {
    'BR': 'Blu-ray',
    'BLU-RAY': 'Blu-ray',
    'CASSETTE': 'Cassette',
    'VINILO': 'Vinilo',
    'LP': 'Vinilo',
}


def limpio(valor):
    if valor is None:
        return None
    texto = ' '.join(str(valor).split())
    return texto or None


def formato(valor):
    texto = limpio(valor) or ''
    return FORMATOS.get(texto.upper(), texto)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        sys.exit('Uso: python scripts/importar-excel.py <Wishlist.xlsx> [--forzar]')
    if DESTINO.exists() and '--forzar' not in sys.argv:
        sys.exit(f'{DESTINO} ya existe. Usa --forzar para sobrescribirlo.')

    libro = openpyxl.load_workbook(args[0], data_only=True)
    hoy = date.today().isoformat()
    items = []

    for fila in libro['Videoteca'].iter_rows(min_row=2, values_only=True):
        titulo, edicion, fmt = (list(fila) + [None] * 3)[:3]
        if not limpio(titulo):
            continue
        items.append({
            'id': uuid.uuid4().hex[:10],
            'seccion': 'videoteca',
            'titulo': limpio(titulo),
            'edicion': limpio(edicion),
            'formato': formato(fmt),
            'anadido': hoy,
        })

    for fila in libro['Discoteca'].iter_rows(min_row=2, values_only=True):
        artista, album, edicion, fmt = (list(fila) + [None] * 4)[:4]
        if not limpio(album):
            continue
        items.append({
            'id': uuid.uuid4().hex[:10],
            'seccion': 'discoteca',
            'artista': limpio(artista),
            'titulo': limpio(album),
            'edicion': limpio(edicion),
            'formato': formato(fmt),
            'anadido': hoy,
        })

    # Quita las claves vacías para que el JSON quede limpio.
    items = [{k: v for k, v in it.items() if v is not None} for it in items]
    DESTINO.parent.mkdir(exist_ok=True)
    DESTINO.write_text(
        json.dumps({'version': 1, 'items': items}, ensure_ascii=False, indent=1) + '\n',
        encoding='utf-8',
    )
    v = sum(1 for i in items if i['seccion'] == 'videoteca')
    print(f'{v} películas y {len(items) - v} discos en {DESTINO}')


if __name__ == '__main__':
    main()
